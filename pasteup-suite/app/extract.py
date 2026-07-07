"""Extract plain text from uploaded DOCX and PDF bill files."""

import io
import os

import fitz  # PyMuPDF
from docx import Document


class ExtractionError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".docx":
        return _from_docx(data)
    if ext == ".pdf":
        return _from_pdf(data)
    if ext in (".txt", ".text"):
        return data.decode("utf-8", errors="replace")
    if ext == ".doc":
        raise ExtractionError(
            "Legacy .doc files are not supported. Open the file in Word and "
            "save it as .docx, then upload again."
        )
    raise ExtractionError(
        f"Unsupported file type '{ext or '(none)'}'. Upload a .docx, .pdf, or .txt file."
    )


def _from_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX file: {exc}") from exc
    parts = []
    for para in doc.paragraphs:
        parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionError(f"Could not read PDF file: {exc}") from exc
    pages = [page.get_text("text") for page in pdf]
    pdf.close()
    return "\n".join(pages)
