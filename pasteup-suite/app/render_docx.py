"""Render an aligned paste-up as a Word document (two-column tables,
differing language highlighted per chamber)."""

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

from .match import Pair
from .wdiff import PARA, Token, diff_units, tokenize

LEFT_HIGHLIGHT = WD_COLOR_INDEX.YELLOW
RIGHT_HIGHLIGHT = WD_COLOR_INDEX.TURQUOISE

STATUS_LABELS = {
    "identical": "IDENTICAL",
    "different": "DIFFERENCES",
    "left_only": "{left} ONLY",
    "right_only": "{right} ONLY",
}


def _write_tokens(paragraph, marked: list[tuple[Token, bool]], highlight) -> None:
    for token, changed in marked:
        if token.text == PARA:
            paragraph.add_run().add_break()
            paragraph.add_run().add_break()
            continue
        if token.space_before:
            paragraph.add_run(" ")
        run = paragraph.add_run(token.text)
        if changed:
            run.font.highlight_color = highlight


def render_docx(pairs: list[Pair], mode: str, left_name: str, right_name: str,
                path: str, left_label: str = "House", right_label: str = "Senate") -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)

    titles = {"standard": "Conference Paste-Up", "bond": "Bond Bill Conference Paste-Up",
              "charter": "Charter Paste-Up"}
    doc.add_heading(titles.get(mode, "Paste-Up"), level=0)
    p = doc.add_paragraph()
    p.add_run(f"{left_label}: {left_name}    {right_label}: {right_name}").italic = True

    counts = {"identical": 0, "different": 0, "left_only": 0, "right_only": 0}
    for pair in pairs:
        counts[pair.status] += 1
    p = doc.add_paragraph()
    p.add_run(
        f"{counts['different']} with differences · {counts['identical']} identical · "
        f"{counts['left_only']} {left_label}-only · {counts['right_only']} {right_label}-only"
    )
    p = doc.add_paragraph()
    r = p.add_run(f"Yellow = {left_label} language not in {right_label}.  ")
    r.font.highlight_color = LEFT_HIGHLIGHT
    r = p.add_run(f"Turquoise = {right_label} language not in {left_label}.")
    r.font.highlight_color = RIGHT_HIGHLIGHT

    for pair in pairs:
        status = pair.status
        unit = pair.left or pair.right
        label = STATUS_LABELS[status].format(left=left_label.upper(), right=right_label.upper())

        if unit.kind == "container" and status == "identical":
            h = doc.add_heading(f"{unit.label} — {label}", level=2)
            continue

        heading = unit.label
        if pair.left and pair.right and pair.left.label != pair.right.label:
            heading = f"{pair.left.label} ↔ {pair.right.label}"
        h = doc.add_heading(f"{heading} — {label}", level=2)
        if status in ("left_only", "right_only"):
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x7B, 0x1F, 0xA2)

        if (pair.left and pair.right and pair.left.kind == "line_item"
                and pair.left.amount is not None and pair.right.amount is not None
                and pair.left.amount != pair.right.amount):
            p = doc.add_paragraph()
            r = p.add_run(
                f"⚑ Amounts differ: {left_label} ${pair.left.amount:,} vs "
                f"{right_label} ${pair.right.amount:,}"
            )
            r.bold = True
            r.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)

        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].paragraphs[0].add_run(left_label).bold = True
        hdr[1].paragraphs[0].add_run(right_label).bold = True

        body = table.rows[1].cells
        if status == "different":
            diff = diff_units(pair.left.text, pair.right.text)
            _write_tokens(body[0].paragraphs[0], diff.left_marked, LEFT_HIGHLIGHT)
            _write_tokens(body[1].paragraphs[0], diff.right_marked, RIGHT_HIGHLIGHT)
        else:
            if pair.left is not None:
                _write_tokens(body[0].paragraphs[0],
                              [(t, False) for t in tokenize(pair.left.text)], None)
            else:
                body[0].paragraphs[0].add_run("— no counterpart —").italic = True
            if pair.right is not None:
                _write_tokens(body[1].paragraphs[0],
                              [(t, False) for t in tokenize(pair.right.text)], None)
            else:
                body[1].paragraphs[0].add_run("— no counterpart —").italic = True

    doc.save(path)
