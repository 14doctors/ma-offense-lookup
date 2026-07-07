import io

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _docx_bytes(text: str) -> bytes:
    doc = Document()
    for para in text.split("\n\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


HOUSE = "An Act.\n\nSECTION 1. Twenty miles per hour.\n\nSECTION 2. Effective January 1."
SENATE = "An Act.\n\nSECTION 1. Twenty-five miles per hour.\n\nSECTION 2. Effective January 1."


def test_healthz():
    assert client.get("/healthz").json() == {"ok": True}


def test_index_renders_form():
    r = client.get("/")
    assert r.status_code == 200
    assert "Generate paste-up" in r.text


def test_full_compare_flow_with_docx_upload():
    r = client.post(
        "/compare",
        files={
            "left": ("house.docx", _docx_bytes(HOUSE)),
            "right": ("senate.docx", _docx_bytes(SENATE)),
        },
        data={"mode": "auto"},
        follow_redirects=True,
    )
    assert r.status_code == 200
    assert "Conference Paste-Up" in r.text
    assert "SECTION 1" in r.text
    assert "diff-right" in r.text          # senate-side highlight present
    assert "Download Word" in r.text

    job_id = r.url.path.rsplit("/", 1)[-1]
    d = client.get(f"/download/{job_id}")
    assert d.status_code == 200
    assert d.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml"
    )
    Document(io.BytesIO(d.content))        # valid docx


def test_unsupported_extension_returns_friendly_error():
    r = client.post(
        "/compare",
        files={
            "left": ("house.rtf", b"x"),
            "right": ("senate.docx", _docx_bytes(SENATE)),
        },
        data={"mode": "auto"},
    )
    assert r.status_code == 422
    assert "Unsupported file type" in r.text


def test_missing_result_404s():
    assert client.get("/result/nope").status_code == 404
